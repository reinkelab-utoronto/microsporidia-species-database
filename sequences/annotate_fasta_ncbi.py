#!/usr/bin/env python3
"""
Annotate FASTA headers with host / locality / date metadata from NCBI Nucleotide,
and emit a spreadsheet laid out in the microsporidia attribute-database column order.

Reads accessions from the headers of a FASTA file, fetches the GenBank records,
pulls qualifiers off the /source feature plus reference and date information,
then writes:

  1. a new FASTA with metadata appended to each header
  2. a full metadata TSV (all extracted fields, one row per accession)
  3. a database-format sheet (.xlsx and .tsv) using your exact column headers,
     ready to paste as new rows into the existing database

Usage
-----
  python annotate_fasta_ncbi.py in.fasta -o out.fasta --email you@utoronto.ca

  # tree-safe labels (no spaces, no punctuation that breaks Newick):
  python annotate_fasta_ncbi.py in.fasta -o out.fasta --email you@utoronto.ca \
      --style tree --label-fields accession,organism,host,geo

Notes
-----
* NCBI requires an email address. An API key (free, from your NCBI account)
  raises the rate limit from 3 to 10 requests/sec -- pass it with --api-key.
* Results are cached in a TSV (--cache, default <output>.meta.tsv) so re-runs
  only fetch accessions that are missing. Delete the cache to force a refetch.
* /country was renamed /geo_loc_name by NCBI in 2023; both are handled.
"""

import argparse
import csv
import os
import re
import sys
import time
from datetime import datetime
from urllib.error import HTTPError, URLError

from Bio import Entrez, SeqIO

# ----------------------------------------------------------------------------
# fields extracted from each GenBank record
# ----------------------------------------------------------------------------
FIELDS = [
    "accession",
    "organism",
    "host",
    "lab_host",
    "geo_loc_name",     # full string, may be "Iran: Tehran"
    "country",          # part before the colon
    "region",           # part after the colon
    "isolate",
    "strain",
    "isolation_source",
    "collection_date",  # /collection_date -- when the sample was taken
    "collection_year",  # 4-digit year parsed out of the above
    "note",
    "lat_lon",
    "pub_year",         # year of the linked publication
    "submitted_date",   # date the record was submitted to GenBank
    "submitted_year",
    "record_date",      # last-modified date on the LOCUS line
    "reference",        # short citation string
    "pubmed",           # PMID(s), comma separated
    "length",
]

# our field name -> GenBank qualifier keys to try, in order
QUALIFIER_ALIASES = {
    "host": ["host", "specific_host"],
    "lab_host": ["lab_host"],
    "geo_loc_name": ["geo_loc_name", "country"],
    "isolate": ["isolate"],
    "strain": ["strain"],
    "isolation_source": ["isolation_source", "tissue_type", "isolation_source"],
    "collection_date": ["collection_date"],
    "note": ["note"],
    "lat_lon": ["lat_lon"],
}

# matches things like MH020212.1, AF123456, NC_002000.2, U10342
ACC_RE = re.compile(r"^([A-Z]{1,6}_?\d{5,9}(?:\.\d+)?)$", re.IGNORECASE)

# ----------------------------------------------------------------------------
# target database column order (exactly as given -- do not reorder)
# ----------------------------------------------------------------------------
DB_COLUMNS = [
    "Timestamp",
    "Species Name",
    "Date Identified (year)",
    "Natural Host(s)",
    "Experimental Host(s)",
    "Host Environment",
    "Host Life Stage during Infection",
    "Site of Infection",
    "Transmission",
    "Spore Length Average (\u00b5m)",
    "Spore Width Average (\u00b5m)",
    "Calculated Volume (\u00b5m\u00b3) (see methods)",
    "Spore Shape (Class; Condition)",
    "Locality",
    "Nucleus",
    "Measured Polar Tubule Length Max (\u03bcm)",
    "Measured Polar Tubule Length Min (\u03bcm)",
    "Measured Polar Tubule (\u03bcm)",
    "Calculated Polar Tubule (\u03bcm) (see methods)",
    "Polar Tubule Coils Range",
    "Polar Tubule Coils Average",
    "18S Accession #",
    "Has the genome been sequenced?",
    "Important Remarks",
    "References",
]


# ----------------------------------------------------------------------------
# FASTA header handling
# ----------------------------------------------------------------------------
def parse_accession(header):
    """Pull a GenBank accession out of a FASTA header line (without '>')."""
    token = header.split()[0] if header.split() else header

    # gi|123456|gb|AF123456.1|  or  lcl|AF123456  or  ref|NC_002000.2|
    if "|" in token:
        parts = [p for p in token.split("|") if p]
        for p in reversed(parts):
            if ACC_RE.match(p):
                return p
        token = parts[-1] if parts else token

    if ACC_RE.match(token):
        return token

    # last resort: first accession-looking substring anywhere in the header,
    # allowing '_' as a delimiter (e.g. Ent_sp_EA38_FJ790319)
    m = re.search(
        r"(?<![A-Za-z0-9])([A-Z]{1,6}_?\d{5,9}(?:\.\d+)?)(?![A-Za-z0-9])", header
    )
    return m.group(1) if m else token


def strip_version(acc):
    return acc.split(".")[0]


def sanitize(text, tree_safe=False):
    """Clean a qualifier value for use in a header."""
    if not text:
        return ""
    text = " ".join(str(text).split())          # collapse whitespace
    if tree_safe:
        text = re.sub(r"[\s,:;()\[\]'\"|/\\]+", "_", text)
        text = re.sub(r"_+", "_", text).strip("_")
    else:
        text = text.replace("|", "/")           # '|' confuses some parsers
    return text


def year_of(text):
    """Pull a 4-digit year out of a date-ish string (1800-2099)."""
    if not text:
        return ""
    m = re.search(r"\b(1[89]\d{2}|20\d{2})\b", str(text))
    return m.group(1) if m else ""


# ----------------------------------------------------------------------------
# NCBI record parsing
# ----------------------------------------------------------------------------
def extract_references(record, info):
    """Fill reference / pubmed / pub_year / submitted_date from the references."""
    citations, pmids, years = [], [], []

    for ref in record.annotations.get("references", []):
        journal = (ref.journal or "").strip()
        title = (ref.title or "").strip()

        # the "Direct Submission" reference carries the submission date
        if journal.startswith("Submitted") or title == "Direct Submission":
            m = re.search(r"Submitted\s*\((\d{1,2}-[A-Z]{3}-\d{4})\)", journal)
            if m and not info["submitted_date"]:
                info["submitted_date"] = m.group(1)
            continue

        if getattr(ref, "pubmed_id", ""):
            pmids.append(str(ref.pubmed_id))

        year = year_of(journal)
        if year:
            years.append(year)

        first_author = ""
        if ref.authors:
            first_author = ref.authors.split(",")[0].strip()
            if " " in first_author:
                first_author = first_author.split()[0]

        bits = [b for b in (first_author, journal) if b]
        if year and year not in journal:
            bits.insert(1 if first_author else 0, f"({year})")
        if bits:
            citations.append(" ".join(bits))

    info["pubmed"] = ",".join(dict.fromkeys(pmids))
    info["reference"] = "; ".join(dict.fromkeys(citations))
    if years:
        info["pub_year"] = min(years)
    info["submitted_year"] = year_of(info["submitted_date"])
    return info


def extract_source_info(record):
    """Pull the fields of interest out of a Biopython GenBank SeqRecord."""
    info = {f: "" for f in FIELDS}
    info["accession"] = record.id
    info["organism"] = record.annotations.get("organism", "")
    info["record_date"] = record.annotations.get("date", "")
    info["length"] = str(len(record.seq)) if record.seq is not None else ""

    for feat in record.features:
        if feat.type != "source":
            continue
        q = feat.qualifiers
        for field, keys in QUALIFIER_ALIASES.items():
            for key in keys:
                if key in q and q[key]:
                    info[field] = "; ".join(str(v) for v in q[key])
                    break
        break  # only the first source feature

    geo = info["geo_loc_name"]
    if geo:
        if ":" in geo:
            country, region = geo.split(":", 1)
            info["country"] = country.strip()
            info["region"] = region.strip()
        else:
            info["country"] = geo.strip()

    info["collection_year"] = year_of(info["collection_date"])
    extract_references(record, info)
    return info


# ----------------------------------------------------------------------------
# NCBI fetching
# ----------------------------------------------------------------------------
def fetch_batch(accessions, retries=4):
    """efetch a batch of accessions as GenBank, return {accession: info}."""
    out = {}
    for attempt in range(retries):
        try:
            handle = Entrez.efetch(
                db="nuccore",
                id=",".join(accessions),
                rettype="gb",
                retmode="text",
            )
            for record in SeqIO.parse(handle, "genbank"):
                out[strip_version(record.id)] = extract_source_info(record)
            handle.close()
            return out
        except (HTTPError, URLError, IOError, ValueError) as err:
            wait = 3 * (attempt + 1)
            sys.stderr.write(
                f"  fetch failed ({err}); retrying in {wait}s "
                f"[{attempt + 1}/{retries}]\n"
            )
            time.sleep(wait)
    sys.stderr.write(f"  giving up on batch of {len(accessions)}\n")
    return out


def fetch_all(accessions, batch_size=150, delay=0.4):
    """Fetch metadata for a list of accessions, in batches."""
    results = {}
    total = len(accessions)
    for i in range(0, total, batch_size):
        batch = accessions[i:i + batch_size]
        sys.stderr.write(
            f"Fetching {i + 1}-{min(i + batch_size, total)} of {total}...\n"
        )
        results.update(fetch_batch(batch))
        time.sleep(delay)
    return results


# ----------------------------------------------------------------------------
# cache / tables
# ----------------------------------------------------------------------------
def load_cache(path):
    if not path or not os.path.exists(path):
        return {}
    cache = {}
    with open(path, newline="", encoding="utf-8") as fh:
        for row in csv.DictReader(fh, delimiter="\t"):
            key = strip_version(row.get("accession", ""))
            if key:
                cache[key] = {f: row.get(f, "") for f in FIELDS}
    sys.stderr.write(f"Loaded {len(cache)} cached records from {path}\n")
    return cache


def write_table(path, accessions, meta):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELDS, delimiter="\t",
                                extrasaction="ignore")
        writer.writeheader()
        for acc in accessions:
            row = meta.get(strip_version(acc))
            if row is None:
                row = {f: "" for f in FIELDS}
                row["accession"] = acc
            writer.writerow(row)


# ----------------------------------------------------------------------------
# database-format sheet
# ----------------------------------------------------------------------------
def build_remarks(info):
    """Everything worth keeping that has no dedicated database column."""
    bits = []
    for label, key in [
        ("isolate", "isolate"),
        ("strain", "strain"),
        ("isolation source", "isolation_source"),
        ("collection date", "collection_date"),
        ("lat/lon", "lat_lon"),
        ("note", "note"),
        ("submitted", "submitted_date"),
    ]:
        val = sanitize(info.get(key, ""))
        if val:
            bits.append(f"{label}: {val}")
    if info.get("length"):
        bits.append(f"18S length: {info['length']} bp")
    return "; ".join(bits)


def build_references(info):
    ref = sanitize(info.get("reference", ""))
    pmid = info.get("pubmed", "")
    acc = info.get("accession", "")
    parts = []
    if ref:
        parts.append(ref)
    if pmid:
        parts.append("PMID: " + pmid)
    if acc:
        parts.append(f"GenBank {acc}")
    return "; ".join(parts)


def build_db_row(acc, info, timestamp, year_from, site_from_source):
    row = {c: "" for c in DB_COLUMNS}
    row["Timestamp"] = timestamp
    row["Species Name"] = sanitize(info.get("organism", ""))
    row["Natural Host(s)"] = sanitize(info.get("host", ""))
    row["Experimental Host(s)"] = sanitize(info.get("lab_host", ""))
    row["Locality"] = sanitize(info.get("geo_loc_name", ""))
    row["18S Accession #"] = info.get("accession", "") or acc
    row["Important Remarks"] = build_remarks(info)
    row["References"] = build_references(info)

    if site_from_source:
        row["Site of Infection"] = sanitize(info.get("isolation_source", ""))

    year = ""
    if year_from == "publication":
        year = info.get("pub_year", "") or info.get("submitted_year", "")
    elif year_from == "submission":
        year = info.get("submitted_year", "")
    elif year_from == "collection":
        year = info.get("collection_year", "")
    row["Date Identified (year)"] = year

    return row


def write_db_tsv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=DB_COLUMNS, delimiter="\t",
                                extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_db_xlsx(path, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.stderr.write(
            "openpyxl not installed -- skipping .xlsx output "
            "(pip install openpyxl). The .tsv version was still written.\n"
        )
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "New entries"

    body = Font(name="Arial", size=10)
    head = Font(name="Arial", size=10, bold=True)

    ws.append(DB_COLUMNS)
    for cell in ws[1]:
        cell.font = head
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in rows:
        ws.append([row.get(c, "") for c in DB_COLUMNS])

    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = body
            cell.alignment = Alignment(vertical="top")
            # keep accessions, years and dates as text so Excel can't reformat them
            if isinstance(cell.value, str):
                cell.number_format = "@"

    widths = {"Species Name": 28, "Natural Host(s)": 28, "Locality": 22,
              "Important Remarks": 55, "References": 45, "Timestamp": 20}
    for i, col in enumerate(DB_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
    ws.freeze_panes = "A2"

    wb.save(path)
    return True


# ----------------------------------------------------------------------------
# header building
# ----------------------------------------------------------------------------
FIELD_SHORTCUTS = {
    "geo": "geo_loc_name",
    "acc": "accession",
    "date": "collection_date",
    "year": "collection_year",
    "source": "isolation_source",
}


def resolve_field(name):
    return FIELD_SHORTCUTS.get(name, name)


def build_header(orig_header, acc, info, style, label_fields, missing):
    """Return the new header line (without '>')."""
    fields = [resolve_field(f.strip()) for f in label_fields.split(",")]

    if style == "tree":
        parts = []
        for f in fields:
            val = acc if f == "accession" else info.get(f, "")
            val = sanitize(val, tree_safe=True)
            parts.append(val if val else missing)
        while len(parts) > 1 and parts[-1] == missing:
            parts.pop()
        return "_".join(parts)

    # style == "append": keep the original header, tack on key=value pairs
    tags = []
    for f in fields:
        if f == "accession":
            continue
        val = sanitize(info.get(f, ""))
        if val:
            tags.append(f"[{f}={val}]")
        elif missing:
            tags.append(f"[{f}={missing}]")
    return orig_header + (" " + " ".join(tags) if tags else "")


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(
        description="Append NCBI host/locality/date metadata to FASTA headers "
                    "and export a database-format spreadsheet.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("fasta", help="input FASTA file")
    p.add_argument("-o", "--out", help="output FASTA (default: <input>.annot.fasta)")
    p.add_argument("--email", help="email for NCBI Entrez (required by NCBI)")
    p.add_argument("--api-key", help="NCBI API key (optional, raises rate limit)")
    p.add_argument("--table", help="full metadata TSV (default: <out>.meta.tsv)")
    p.add_argument("--cache", help="cache TSV to read/write (default: same as --table)")
    p.add_argument("--db-sheet", help="database-format output, .xlsx "
                                      "(default: <out>.db.xlsx); a .tsv twin is "
                                      "always written alongside it")
    p.add_argument("--style", choices=["append", "tree"], default="append",
                   help="'append' keeps the original header and adds [key=value] "
                        "tags; 'tree' replaces it with an underscore-joined label")
    p.add_argument("--label-fields", default="accession,host,geo",
                   help="comma-separated fields to put in the header. Options: "
                        + ", ".join(FIELDS) + " (shortcuts: geo, acc, date, year, source)")
    p.add_argument("--missing", default="NA",
                   help="placeholder for missing values ('' to omit them)")
    p.add_argument("--timestamp", default=None,
                   help="value for the Timestamp column (default: now)")
    p.add_argument("--year-from", choices=["none", "publication", "submission",
                                           "collection"], default="none",
                   help="what to put in 'Date Identified (year)'. GenBank has no "
                        "species-description year, so the default leaves it blank; "
                        "all three dates are in the metadata TSV and Remarks either way")
    p.add_argument("--site-from-isolation-source", action="store_true",
                   help="copy /isolation_source into 'Site of Infection' "
                        "(off by default -- it often holds a substrate, not a tissue)")
    p.add_argument("--one-row-per-species", action="store_true",
                   help="collapse the database sheet to one row per organism name "
                        "instead of one row per sequence")
    p.add_argument("--batch-size", type=int, default=150)
    p.add_argument("--line-width", type=int, default=60,
                   help="sequence line wrap width; 0 for a single line")
    p.add_argument("--no-fetch", action="store_true",
                   help="use only the cache, don't contact NCBI")
    args = p.parse_args()

    stem = re.sub(r"\.(fa|fasta|fna|fas)$", "", args.fasta)
    out_path = args.out or stem + ".annot.fasta"
    table_path = args.table or out_path + ".meta.tsv"
    cache_path = args.cache or table_path
    db_path = args.db_sheet or out_path + ".db.xlsx"
    db_tsv_path = re.sub(r"\.xlsx$", "", db_path) + ".tsv"
    timestamp = args.timestamp if args.timestamp is not None else \
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # --- read accessions from the FASTA -------------------------------------
    records = list(SeqIO.parse(args.fasta, "fasta"))
    if not records:
        sys.exit(f"No FASTA records found in {args.fasta}")
    headers = [r.description for r in records]
    accessions = [parse_accession(h) for h in headers]
    sys.stderr.write(f"Read {len(records)} sequences from {args.fasta}\n")

    # --- fetch what we don't already have -----------------------------------
    meta = load_cache(cache_path)
    needed = sorted({strip_version(a) for a in accessions} - set(meta))

    if needed and not args.no_fetch:
        if not args.email:
            sys.exit("--email is required to query NCBI (or use --no-fetch)")
        Entrez.email = args.email
        if args.api_key:
            Entrez.api_key = args.api_key
        delay = 0.15 if args.api_key else 0.4
        fetched = fetch_all(needed, batch_size=args.batch_size, delay=delay)
        meta.update(fetched)
        missing_accs = [a for a in needed if a not in fetched]
        if missing_accs:
            sys.stderr.write(
                f"WARNING: no record returned for {len(missing_accs)} accession(s): "
                f"{', '.join(missing_accs[:10])}"
                f"{' ...' if len(missing_accs) > 10 else ''}\n"
            )
    elif needed:
        sys.stderr.write(f"{len(needed)} accessions not in cache (--no-fetch set)\n")

    write_table(table_path, accessions, meta)
    if cache_path != table_path:
        write_table(cache_path, sorted(meta), meta)

    # --- database-format sheet ----------------------------------------------
    db_rows, seen_species = [], {}
    for acc in accessions:
        info = meta.get(strip_version(acc))
        if info is None:
            info = {f: "" for f in FIELDS}
            info["accession"] = acc
        row = build_db_row(acc, info, timestamp, args.year_from,
                           args.site_from_isolation_source)

        if args.one_row_per_species:
            key = row["Species Name"] or acc
            if key in seen_species:
                prev = seen_species[key]
                for col in ("Natural Host(s)", "Locality", "18S Accession #"):
                    vals = [v for v in (prev[col], row[col]) if v]
                    prev[col] = "; ".join(dict.fromkeys(
                        v for part in vals for v in part.split("; ")))
                continue
            seen_species[key] = row
        db_rows.append(row)

    write_db_tsv(db_tsv_path, db_rows)
    wrote_xlsx = write_db_xlsx(db_path, db_rows)

    # --- write the annotated FASTA ------------------------------------------
    seen, n_host, n_geo, n_date = {}, 0, 0, 0
    with open(out_path, "w") as out:
        for rec, header, acc in zip(records, headers, accessions):
            info = meta.get(strip_version(acc), {f: "" for f in FIELDS})
            n_host += bool(info.get("host"))
            n_geo += bool(info.get("geo_loc_name"))
            n_date += bool(info.get("collection_date"))

            new_header = build_header(header, acc, info, args.style,
                                      args.label_fields, args.missing)

            # keep labels unique (important for tree style)
            if new_header in seen:
                seen[new_header] += 1
                new_header = f"{new_header}_{seen[new_header]}"
            else:
                seen[new_header] = 0

            seq = str(rec.seq)
            out.write(f">{new_header}\n")
            if args.line_width and args.line_width > 0:
                for i in range(0, len(seq), args.line_width):
                    out.write(seq[i:i + args.line_width] + "\n")
            else:
                out.write(seq + "\n")

    n = len(records)
    sys.stderr.write(
        f"\nWrote {n} sequences to {out_path}\n"
        f"Metadata table:  {table_path}\n"
        f"Database sheet:  {db_tsv_path}"
        + (f" and {db_path}\n" if wrote_xlsx else "\n")
        + f"Recovered host {n_host}/{n}, locality {n_geo}/{n}, "
          f"collection date {n_date}/{n}\n"
    )


if __name__ == "__main__":
    main()
